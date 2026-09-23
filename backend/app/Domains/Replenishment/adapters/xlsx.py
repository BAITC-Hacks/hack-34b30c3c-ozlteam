"""Seven-sheet 1C-friendly workbook interchange using the existing openpyxl dependency."""

from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError

from app.Domains.Replenishment.DTO.calculation import CalculationInput

SHEETS = {
    "products": ("sku", "name", "category", "supplier_id"),
    "suppliers": ("id", "name", "lead_days", "min_order_qty"),
    "sales": ("date", "sku", "warehouse", "customer_id", "quantity", "price"),
    "stock": ("sku", "warehouse", "quantity"),
    "inbound": ("sku", "warehouse", "quantity", "eta"),
    "stockouts": ("sku", "warehouse", "start", "end"),
    "category_growth": ("category", "growth_pct"),
}
SHEET_LABELS = {
    "products": "Товары",
    "suppliers": "Поставщики",
    "sales": "Продажи",
    "stock": "Остатки",
    "inbound": "Товары в пути",
    "stockouts": "Периоды отсутствия",
    "category_growth": "Прирост по категориям",
    "config": "Настройки",
}
COLUMN_LABELS = {
    "sku": "Артикул",
    "name": "Наименование",
    "category": "Категория",
    "supplier_id": "Код поставщика",
    "id": "Код поставщика",
    "lead_days": "Срок поставки, дней",
    "min_order_qty": "Минимальная партия",
    "date": "Дата продажи",
    "warehouse": "Склад",
    "customer_id": "Обезличенный код клиента",
    "quantity": "Количество",
    "price": "Цена за единицу",
    "eta": "Ожидаемая дата прихода",
    "start": "Начало отсутствия",
    "end": "Конец отсутствия",
    "growth_pct": "Ожидаемый прирост, %",
    "setting": "Параметр",
    "value": "Значение",
}
SETTING_LABELS = {"as_of": "Дата расчёта", "review_days": "Период пересмотра, дней"}
MAX_ROWS_PER_SHEET = 100_000
MAX_UNCOMPRESSED_BYTES = 150 * 1024 * 1024


class InvalidWorkbook(ValueError):
    pass


def parse_workbook(content: bytes) -> CalculationInput:
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise InvalidWorkbook("Распакованный XLSX превышает 150 МиБ")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, OSError, ValueError) as error:
        raise InvalidWorkbook("Файл не является корректной книгой XLSX") from error
    try:
        missing = [
            SHEET_LABELS[name]
            for name in SHEETS
            if name not in workbook and SHEET_LABELS[name] not in workbook
        ]
        if missing:
            raise InvalidWorkbook(f"Отсутствуют листы: {', '.join(sorted(missing))}")
        payload: dict = {}
        for name, columns in SHEETS.items():
            sheet_name = SHEET_LABELS[name] if SHEET_LABELS[name] in workbook else name
            if name in workbook and SHEET_LABELS[name] in workbook:
                raise InvalidWorkbook(f"Листы {name} и {SHEET_LABELS[name]} дублируют друг друга")
            rows = workbook[sheet_name].iter_rows(values_only=True)
            header = next(rows, None)
            labels = tuple(COLUMN_LABELS[column] for column in columns)
            if header is None or tuple(header[: len(columns)]) not in (columns, labels):
                raise InvalidWorkbook(f"Лист {sheet_name}: ожидаются колонки {', '.join(labels)}")
            parsed = []
            for number, cells in enumerate(rows, start=2):
                if number > MAX_ROWS_PER_SHEET + 1:
                    raise InvalidWorkbook(f"Лист {sheet_name}: максимум {MAX_ROWS_PER_SHEET} строк")
                if all(value is None for value in cells):
                    continue
                if any(cells[index] is None for index in range(len(columns))):
                    raise InvalidWorkbook(
                        f"Лист {sheet_name}, строка {number}: пустое обязательное поле"
                    )
                row = dict(zip(columns, cells, strict=False))
                for field in ("date", "eta", "start", "end"):
                    if field in row and isinstance(row[field], datetime):
                        row[field] = row[field].date()
                    elif field in row and isinstance(row[field], str):
                        try:
                            row[field] = date.fromisoformat(row[field])
                        except ValueError as error:
                            raise InvalidWorkbook(
                                f"Лист {sheet_name}, строка {number}: "
                                f"неверная дата в {COLUMN_LABELS[field]}"
                            ) from error
                parsed.append(row)
            payload[name] = parsed
        if "config" in workbook or SHEET_LABELS["config"] in workbook:
            if "config" in workbook and SHEET_LABELS["config"] in workbook:
                raise InvalidWorkbook("Листы config и Настройки дублируют друг друга")
            config_name = SHEET_LABELS["config"] if SHEET_LABELS["config"] in workbook else "config"
            rows = workbook[config_name].iter_rows(values_only=True)
            if next(rows, None) not in (("setting", "value"), ("Параметр", "Значение")):
                raise InvalidWorkbook("Лист Настройки: ожидаются колонки Параметр, Значение")
            setting_keys = {label: key for key, label in SETTING_LABELS.items()}
            settings = {
                setting_keys.get(cells[0], cells[0]): cells[1]
                for cells in rows
                if cells and cells[0] is not None
            }
            if "as_of" not in settings or "review_days" not in settings:
                raise InvalidWorkbook("Лист Настройки: нужны дата расчёта и период пересмотра")
            payload["as_of"] = settings["as_of"]
            payload["review_days"] = settings["review_days"]
            if isinstance(payload["as_of"], datetime):
                payload["as_of"] = payload["as_of"].date()
        else:
            if not payload["sales"]:
                raise InvalidWorkbook("Без листа config нужна хотя бы одна дата продажи")
            payload["as_of"] = max(row["date"] for row in payload["sales"])
            payload["review_days"] = 14
        try:
            return CalculationInput.model_validate(payload)
        except ValidationError as error:
            first = error.errors()[0]
            path = ".".join(str(part) for part in first["loc"])
            raise InvalidWorkbook(f"Некорректные данные {path}: {first['msg']}") from error
    finally:
        workbook.close()


def create_workbook(data: CalculationInput) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    supplier_names = {supplier.id: supplier.name for supplier in data.suppliers}
    for name, columns in SHEETS.items():
        sheet = workbook.create_sheet(SHEET_LABELS[name])
        headers = [COLUMN_LABELS[column] for column in columns]
        if name == "products":
            headers.append("Поставщик (справочно)")
        sheet.append(headers)
        for record in getattr(data, name):
            values = [getattr(record, column) for column in columns]
            if name == "products":
                values.append(supplier_names[record.supplier_id])
            sheet.append(values)
        if name == "products":
            sheet.column_dimensions["D"].hidden = True
        if name == "category_growth":
            for cell in sheet["B"][1:]:
                cell.number_format = "0.0%"
    config = workbook.create_sheet(SHEET_LABELS["config"])
    config.append(("Параметр", "Значение"))
    config.append((SETTING_LABELS["as_of"], data.as_of))
    config.append((SETTING_LABELS["review_days"], data.review_days))
    instructions = workbook.create_sheet("Как заполнить")
    instructions.append(("Демонстрационный шаблон", "Как работать с файлом"))
    instructions.append(
        ("Пример", "Синтетические данные. Это не заказ поставщику и не выгрузка 1С.")
    )
    instructions.append(
        ("Заполнение", "Замените примеры своими данными. Сохраните заголовки и порядок колонок.")
    )
    instructions.append(
        (
            "Поставщик",
            "Название в листе «Товары» дано справочно. Связь задаёт код в скрытой колонке D; "
            "при смене поставщика раскройте её и укажите код из листа «Поставщики».",
        )
    )
    instructions.append(
        (
            "Количество",
            "Указывайте количество в согласованной базовой единице товара. "
            "Артикулы должны совпадать на всех листах.",
        )
    )
    instructions.append(("Прирост", "Вводите процент, например 8%. В расчёт передаётся 0,08."))
    instructions.append(
        (
            "Периоды отсутствия",
            "Указывайте только подтверждённые интервалы отсутствия товара. "
            "Если их нет, оставьте лист пустым с заголовками.",
        )
    )
    instructions.append(
        (
            "Загрузка",
            "Все семь листов данных обязательны. "
            "Настройки задают дату расчёта и период пересмотра.",
        )
    )
    for sheet in workbook:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 34
        sheet.print_title_rows = "1:1"
        for column in sheet.columns:
            letter = column[0].column_letter
            longest = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[letter].width = min(54, max(20, longest + 3))
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="224360")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 34 if sheet != instructions else 78
            for cell in row:
                cell.font = Font(name="Calibri", size=11, color="202B33")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F2F5F8")
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                if isinstance(cell.value, date):
                    cell.number_format = "dd.mm.yyyy"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_options.horizontalCentered = True
        sheet.print_area = sheet.dimensions
    instructions.column_dimensions["B"].width = 90
    instructions.auto_filter.ref = None
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
