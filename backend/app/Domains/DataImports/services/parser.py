import csv
import io
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from pydantic import ValidationError

from app.core.errors import DomainError
from app.Domains.DataImports.DTO.rows import KINDS, ROW_ADAPTER, json_value

MAX_ROWS = 10000
MAX_BYTES = 25 * 1024 * 1024
MAX_UNCOMPRESSED = 150 * 1024 * 1024


def parse_file(
    filename: str, content: bytes, kind: str, mapping: dict[str, str], quantity_multiplier: int = 1
):
    """Bounded normalized import. Never store unselected columns or raw validation inputs."""
    if kind not in KINDS:
        raise DomainError("Неизвестный вид данных", status_code=422, code="invalid_import_kind")
    if len(content) > MAX_BYTES:
        raise DomainError("Файл превышает 25 МиБ", status_code=413, code="file_too_large")
    if quantity_multiplier not in (-1, 1):
        raise DomainError("Множитель количества: 1 или -1", status_code=422, code="invalid_sign")
    if len(set(mapping.values())) != len(mapping):
        raise DomainError(
            "Несколько колонок сопоставлены одному полю", status_code=422, code="invalid_mapping"
        )
    suffix = Path(filename).suffix.lower()
    book = None
    try:
        if suffix == ".csv":
            stream = io.StringIO(content.decode("utf-8-sig"))
            sample = stream.read(4096)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            tables = [("CSV", iter(csv.reader(stream, dialect)))]
        elif suffix == ".xlsx":
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > MAX_UNCOMPRESSED:
                    raise DomainError(
                        "Распакованный XLSX превышает 150 МиБ",
                        status_code=413,
                        code="file_too_large",
                    )
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            tables = [(sheet.title, sheet.iter_rows(values_only=True)) for sheet in book.worksheets]
        else:
            raise DomainError(
                "Поддерживаются CSV UTF-8 и XLSX", status_code=415, code="unsupported_file_type"
            )
        rows, errors, count = [], [], 0
        for sheet, iterator in tables:
            header_row = next(iterator, None)
            if header_row is None:
                continue
            headers = [str(v).strip() if v is not None else "" for v in header_row]
            nonempty = [h for h in headers if h]
            if len(nonempty) != len(set(nonempty)):
                errors.append(dict(sheet=sheet, row=1, column=None, message="Повтор колонок"))
                continue
            missing = set(mapping) - set(headers)
            if missing:
                errors.append(
                    dict(
                        sheet=sheet,
                        row=1,
                        column=None,
                        message="В файле отсутствуют колонки из сопоставления",
                    )
                )
                continue
            for row_number, cells in enumerate(iterator, 2):
                if all(v is None or v == "" for v in cells):
                    continue
                count += 1
                if count > MAX_ROWS:
                    raise DomainError(
                        "Не более 10000 строк в пакете; разделите выгрузку",
                        status_code=413,
                        code="too_many_rows",
                    )
                record = {"kind": kind}
                for key, value in zip(headers, cells, strict=False):
                    if not key or (mapping and key not in mapping):
                        continue
                    target = mapping.get(key, key)
                    if value is not None and value != "":
                        record[target] = json_value(value)
                if record.get("kind") != kind:
                    errors.append(
                        dict(
                            sheet=sheet,
                            row=row_number,
                            column="kind",
                            message="Вид строки не соответствует выбранному виду данных",
                        )
                    )
                    continue
                try:
                    if "quantity" in record and quantity_multiplier == -1:
                        record["quantity"] = str(-Decimal(str(record["quantity"])))
                    validated = ROW_ADAPTER.validate_python(record)
                    rows.append(
                        dict(sheet=sheet, row=row_number, data=validated.model_dump(mode="json"))
                    )
                except (ValidationError, InvalidOperation) as exc:
                    if isinstance(exc, ValidationError):
                        for error in exc.errors(include_input=False, include_context=False):
                            errors.append(
                                dict(
                                    sheet=sheet,
                                    row=row_number,
                                    column=str(error["loc"][-1]) if error["loc"] else None,
                                    message=error["msg"],
                                )
                            )
                    else:
                        errors.append(
                            dict(
                                sheet=sheet,
                                row=row_number,
                                column="quantity",
                                message="Количество не является числом",
                            )
                        )
        if count == 0:
            errors.append(dict(sheet=None, row=0, column=None, message="Файл не содержит данных"))
        return rows, errors, count
    except (UnicodeDecodeError, BadZipFile, csv.Error, ValueError, KeyError) as exc:
        raise DomainError(
            "Не удалось прочитать файл; проверьте формат и кодировку",
            status_code=422,
            code="invalid_file",
        ) from exc
    finally:
        if book is not None:
            book.close()
