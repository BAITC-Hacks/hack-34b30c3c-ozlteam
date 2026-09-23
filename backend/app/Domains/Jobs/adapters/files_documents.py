from asyncio import to_thread
from io import BytesIO
from uuid import UUID

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from app.Domains.Files.contracts import FilesError
from app.Domains.Files.services.file_service import FileService
from app.Domains.Jobs.contracts import Document, DocumentsUnavailable

# Long documents are truncated: a model call must stay predictable in cost and latency.
MAX_DOCUMENT_CHARACTERS = 30_000

IMAGE = "image"
PDF = "pdf"
DOC = "doc"
SHEET = "sheet"


class FilesDocumentSource:
    """Reads an uploaded file through the Files domain and turns it into plain text."""

    def __init__(self, files: FileService, max_characters: int = MAX_DOCUMENT_CHARACTERS):
        self.files = files
        self.max_characters = max_characters

    async def read(self, file_id: UUID) -> Document:
        try:
            file = await self.files.get(file_id)
            download = await self.files.download(file_id)
            data = b"".join([chunk async for chunk in download.stream])
        except FilesError as error:
            raise DocumentsUnavailable(f"File {file_id} is not readable: {error}") from error
        if file.kind == IMAGE:
            return Document(name=file.filename, text="", images=[data])
        if file.filename.lower().endswith(".csv"):
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError as error:
                raise DocumentsUnavailable("CSV must be encoded as UTF-8") from error
            return Document(name=file.filename, text=text[: self.max_characters], images=[])
        # Parsing is CPU bound and synchronous; keep it off the worker event loop.
        text = await to_thread(self._extract, file.kind, data)
        return Document(name=file.filename, text=text[: self.max_characters], images=[])

    def _extract(self, kind: str, data: bytes) -> str:
        if kind == PDF:
            return self._from_pdf(data)
        if kind == DOC:
            return self._from_docx(data)
        if kind == SHEET:
            return self._from_xlsx(data)
        raise DocumentsUnavailable(f"Cannot extract text from files of kind '{kind}'")

    @staticmethod
    def _from_pdf(data: bytes) -> str:
        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()

    @staticmethod
    def _from_docx(data: bytes) -> str:
        document = DocxDocument(BytesIO(data))
        lines: list[str] = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            lines.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(line for line in lines if line.strip()).strip()

    @staticmethod
    def _from_xlsx(data: bytes) -> str:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(f"# {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    cells = ["" if value is None else str(value) for value in row]
                    if any(cell.strip() for cell in cells):
                        lines.append("\t".join(cells))
            return "\n".join(lines).strip()
        finally:
            workbook.close()
