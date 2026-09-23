"""Readable documents from the immutable approved order; exchange stays separate."""

import csv
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from math import ceil
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from app.Domains.Procurement.resources.order import OrderOut

HEADERS = ["№", "Артикул", "Код 1С", "Наименование", "Ед.", "Количество"]
WIDTHS = [6, 24, 20, 68, 10, 25]


def quantity_text(quantity: Decimal) -> str:
    value = format(quantity, "f")
    return value.rstrip("0").rstrip(".") if "." in value else value


def local_date(value: datetime | None) -> str:
    if value is None:
        return "Не указана"
    # Historical naive timestamps in snapshots are stored as UTC.
    return (value.replace(tzinfo=UTC) if value.tzinfo is None else value).astimezone(
        ZoneInfo("Asia/Almaty")
    ).strftime("%d.%m.%Y %H:%M") + " (Алматы)"


def safe_csv(value: str) -> str:
    if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
        return "'" + value
    return value


def line_values(order: OrderOut):
    codes = {row.id: row.code for row in order.external_references.products}
    for index, line in enumerate(order.lines, 1):
        yield [
            index,
            line.sku,
            codes.get(line.product_id) or "",
            line.name,
            line.unit,
            line.quantity,
        ]


def render_order_document(order: OrderOut, file_format: str) -> bytes:
    if file_format == "csv":
        return _csv(order)
    return _xlsx(order)


def _csv(order: OrderOut) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=";")
    writer.writerow(
        [
            "Заказ (ссылка)",
            "Редакция",
            "Дата создания",
            "Дата утверждения",
            "Поставщик",
            "Склад",
            "Комментарий",
            *HEADERS,
            "Идентификатор заказа",
        ]
    )
    for row in line_values(order):
        row[-1] = quantity_text(row[-1]).replace(".", ",")
        values = [
            str(order.id)[-8:].upper(),
            order.revision,
            local_date(order.created_at),
            local_date(order.approved_at),
            order.supplier_name,
            order.external_references.warehouse.name,
            order.comment,
            *row,
            str(order.id),
        ]
        writer.writerow([safe_csv(str(value)) for value in values])
    return output.getvalue().encode("utf-8-sig")


def _text(cell, value: str):
    cell.value = value
    cell.data_type = "s"
    cell.number_format = "@"


def _quantity(cell, value: Decimal):
    text = quantity_text(value)
    significant = len(value.normalize().as_tuple().digits)
    numeric = float(value)
    if significant <= 15 and Decimal(str(numeric)) == value:
        cell.value = numeric
        cell.number_format = "0" if value == value.to_integral_value() else "0.######"
    else:
        _text(cell, text)
    cell.alignment = Alignment(horizontal="right", vertical="top")


def _xlsx(order: OrderOut) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заказ"
    sheet.sheet_view.showGridLines = False
    for index, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    def full_row(value: str, *, bold=False, size=11):
        row = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = sheet.cell(row, 1)
        _text(cell, value)
        cell.font = Font(name="Calibri", size=size, bold=bold, color="202B3A")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = max(
            23, 17 * sum(max(1, ceil(len(part) / 140)) for part in value.split("\n"))
        )
        return row

    full_row("Заказ поставщику", bold=True, size=20)
    full_row(f"Ссылка на заказ: {str(order.id)[-8:].upper()} · Редакция {order.revision}")
    full_row(f"Поставщик: {order.supplier_name}", bold=True)
    full_row(f"Склад: {order.external_references.warehouse.name}")
    full_row(f"Создан: {local_date(order.created_at)} · Утверждён: {local_date(order.approved_at)}")
    if order.comment:
        full_row(f"Комментарий: {order.comment}")
    full_row("Ссылка на заказ — внутреннее обозначение, не номер документа 1С.", size=10)
    sheet.append([])
    sheet.append(HEADERS)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="24374B")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 29
    border = Border(bottom=Side(style="hair", color="DCE2E8"))
    for values in line_values(order):
        sheet.append([None] * 6)
        row = sheet.max_row
        # append(None) still advances the row; all source strings are explicitly text.
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row, column)
            if column == 1:
                cell.value = value
            elif column == 6:
                _quantity(cell, value)
            else:
                _text(cell, value)
            cell.font = Font(name="Calibri", size=11, color="202B3A")
            cell.border = border
            cell.alignment = Alignment(
                wrap_text=True, vertical="top", horizontal="right" if column == 6 else "left"
            )
            if (row - header_row) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F5F8")
        height = max(
            sum(
                max(1, ceil(len(part) / max(1, WIDTHS[column] - 3)))
                for part in str(value).split("\n")
            )
            for column, value in enumerate(values)
        )
        sheet.row_dimensions[row].height = max(29, 16 * height + 10)
    last_line = sheet.max_row
    sheet.auto_filter.ref = f"A{header_row}:F{last_line}"
    sheet.freeze_panes = f"D{header_row + 1}"
    sheet.print_title_rows = f"{header_row}:{header_row}"
    full_row(f"Всего позиций: {len(order.lines)}", bold=True)
    full_row("Утверждено в системе. Отправка поставщику выполняется отдельно.", size=10)
    full_row(f"Идентификатор заказа: {order.id}", size=9)
    sheet.print_area = f"A1:F{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4)
    sheet.oddFooter.right.text = "Страница &P из &N"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
