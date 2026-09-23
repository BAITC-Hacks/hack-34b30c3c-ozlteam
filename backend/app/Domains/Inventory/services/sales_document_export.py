"""Readable export of stored document lines, never an original source invoice."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_DOCUMENT_LINES = 100_000
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def create_sales_document(source_id, document_id, rows) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Продажи"
    sheet.append(["Артикул", "Наименование", "Код 1С", "Склад", "Кол-во", "Ед.", "Статус"])
    for sale, product, warehouse in rows:
        sheet.append(
            [
                product.sku,
                product.name,
                product.code,
                warehouse.name,
                format(sale.quantity, "f"),
                product.unit,
                {"posted": "Проведена", "cancelled": "Отменена"}.get(sale.status, sale.status),
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    for index, width in enumerate([22, 54, 22, 24, 24, 10, 16], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        row[4].alignment = Alignment(horizontal="right", vertical="top")
    info = workbook.create_sheet("О выгрузке")
    info.append(["Параметр", "Значение"])
    info.append(["Документ", document_id])
    info.append(["Источник", str(source_id)])
    info.append(["Описание", "Выгрузка загруженных строк документа. Не оригинал накладной."])
    info.append(["Строк", len(rows)])
    info.append(
        ["Состав", "Все загруженные строки этого документа и источника, включая отменённые."]
    )
    info.append(
        [
            "Количество",
            "Сохранено текстом без округления, чтобы не терять точность исходных данных.",
        ]
    )
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 90
    for tab in workbook:
        tab.sheet_view.showGridLines = False
        tab.row_dimensions[1].height = 28
        for cell in tab[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="224360")
        for row in tab.iter_rows(min_row=2):
            tab.row_dimensions[row[0].row].height = 42
            for cell in row:
                cell.font = Font(name="Calibri", size=11, color="202B33")
                cell.alignment = Alignment(
                    horizontal=cell.alignment.horizontal, vertical="top", wrap_text=True
                )
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F2F5F8")
                if isinstance(cell.value, str):
                    # Treat every external string literally, including leading '=' and '+';
                    # exact Decimal strings also avoid Excel's 15-digit numeric limit.
                    cell.data_type = "s"
                    cell.number_format = "@"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
