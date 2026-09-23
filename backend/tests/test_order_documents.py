import csv
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.core.errors import DomainError
from app.Domains.Procurement.adapters.order_document import HEADERS, render_order_document
from app.Domains.Procurement.resources.order import OrderOut
from app.Domains.Procurement.services.order_service import OrderService


@pytest.fixture
def approved_order():
    supplier, warehouse, source, actor = (uuid4() for _ in range(4))
    product_ids = [uuid4() for _ in range(3)]

    def reference(identifier, name):
        return dict(id=identifier, source_id=source, external_id=str(identifier), name=name)

    return OrderOut(
        id=uuid4(),
        supplier_id=supplier,
        supplier_name='IEK; "Тест"',
        warehouse_id=warehouse,
        status="approved",
        version=2,
        revision=1,
        comment="[ТЕСТОВЫЙ ЗАКАЗ]\nПроверка",
        created_by=actor,
        created_at=datetime(2026, 9, 23, 20, 0, tzinfo=UTC),
        approved_by=actor,
        approved_at=datetime(2026, 9, 23, 21, 0, tzinfo=UTC),
        external_references=dict(
            supplier=reference(supplier, "IEK"),
            warehouse=reference(warehouse, "Алматы"),
            products=[
                dict(**reference(product, "Кабель"), sku="00012", code="00007", unit="шт")
                for product in product_ids
            ],
        ),
        lines=[
            dict(
                id=uuid4(),
                recommendation_id=None,
                run_id=None,
                product_id=product,
                sku=["00012", "=SUM(1,2)", "@danger"][index],
                name=['Кабель; "силовой" ' * 15, '=HYPERLINK("https://example.com")', "+товар"][
                    index
                ],
                unit="шт" if index != 1 else "м",
                recommended_quantity=None,
                quantity=[
                    Decimal("1.000000"),
                    Decimal("123.450000"),
                    Decimal("12345678901234.123456"),
                ][index],
                reason="Тест",
            )
            for index, product in enumerate(product_ids)
        ],
    )


def test_document_xlsx_readable_precision_layout_and_safe_text(approved_order):
    workbook = load_workbook(BytesIO(render_order_document(approved_order, "xlsx")))
    sheet = workbook["Заказ"]
    values = [cell.value for row in sheet for cell in row]
    assert "Заказ поставщику" in values
    assert 'Поставщик: IEK; "Тест"' in values
    assert "Склад: Алматы" in values
    assert any("24.09.2026 01:00" in value for value in values if isinstance(value, str))
    assert any("[ТЕСТОВЫЙ ЗАКАЗ]" in value for value in values if isinstance(value, str))
    header = next(row[0].row for row in sheet if row[0].value == "№")
    assert [cell.value for cell in sheet[header]] == HEADERS
    assert sheet.cell(header, 1).fill.fgColor.rgb == "0024374B"
    assert sheet.cell(header, 1).font.bold
    first = header + 1
    assert sheet.cell(first, 2).value == "00012"
    assert sheet.cell(first, 2).data_type == "s"
    assert sheet.cell(first, 3).value == "00007"
    assert sheet.cell(first, 4).value == approved_order.lines[0].name
    assert sheet.cell(first, 4).alignment.wrap_text
    assert sheet.row_dimensions[first].height > 29
    assert sheet.cell(first, 6).value == 1
    assert sheet.cell(first + 1, 6).value == 123.45
    assert sheet.cell(first + 2, 6).value == "12345678901234.123456"
    assert sheet.cell(first + 2, 6).data_type == "s"
    assert sheet.cell(first + 1, 2).value == "=SUM(1,2)"
    assert all(cell.data_type != "f" for row in sheet for cell in row)
    assert sum(str(approved_order.id) in str(value) for value in values) == 1
    assert str(approved_order.supplier_id) not in values
    assert sheet.freeze_panes == f"D{first}"
    assert sheet.auto_filter.ref == f"A{header}:F{first + 2}"
    assert sheet.print_title_rows == f"${header}:${header}"
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.page_setup.fitToHeight == 0
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.column_dimensions["D"].width >= 60
    workbook.close()


def test_document_csv_russian_headers_exact_quantity_and_quoting(approved_order):
    content = render_order_document(approved_order, "csv")
    assert content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(StringIO(content.decode("utf-8-sig")), delimiter=";"))
    assert len(rows) == 3
    assert rows[0]["Поставщик"] == approved_order.supplier_name
    assert rows[0]["Склад"] == "Алматы"
    assert rows[0]["Артикул"] == "00012"
    assert rows[0]["Код 1С"] == "00007"
    assert rows[0]["Наименование"] == approved_order.lines[0].name
    assert rows[0]["Комментарий"] == approved_order.comment
    assert [row["Количество"] for row in rows] == ["1", "123,45", "12345678901234,123456"]
    assert rows[1]["Артикул"] == "'=SUM(1,2)"
    assert rows[1]["Наименование"].startswith("'=")
    assert rows[2]["Артикул"] == "'@danger"
    assert rows[2]["Наименование"] == "'+товар"
    assert rows[0]["Идентификатор заказа"] == str(approved_order.id)


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
async def test_document_uses_approved_snapshot_and_preserves_exchange(approved_order, file_format):
    stored = SimpleNamespace(
        approved_snapshot=approved_order.model_dump(mode="json"),
        deleted_at=None,
        supplier_name="Changed catalog",
        comment="Changed after approval",
    )

    class Repository:
        async def get(self, identifier, lock=False):
            return stored

    service = OrderService(Repository(), None)
    before = await service.export(approved_order.id, file_format)
    document = await service.export(approved_order.id, file_format, "document")
    expected_document = render_order_document(approved_order, file_format)
    explicit_exchange = await service.export(approved_order.id, file_format, "exchange")

    def contents(content):
        if file_format == "csv":
            return content
        workbook = load_workbook(BytesIO(content))
        values = list(workbook.active.values)
        workbook.close()
        return values

    # ZIP container timestamps do not define document contents.
    assert contents(document) == contents(expected_document)
    assert contents(before) == contents(explicit_exchange)
    if file_format == "csv":
        rows = list(csv.reader(StringIO(before.decode("utf-8-sig"))))
        assert rows[0] == [
            "order_id",
            "revision",
            "supplier_id",
            "warehouse_id",
            "sku",
            "name",
            "unit",
            "quantity",
        ]
        assert rows[1][-1] == "1.000000"
        assert rows[3][-1] == "12345678901234.123456"
    else:
        workbook = load_workbook(BytesIO(before))
        assert workbook.active.title == "Order"
        assert workbook.active["H2"].value == "1.000000"
        assert workbook.active["H4"].value == "12345678901234.123456"
        workbook.close()


async def test_document_draft_export_still_requires_approval(approved_order):
    class Repository:
        async def get(self, identifier, lock=False):
            return SimpleNamespace(
                approved_snapshot=approved_order.model_copy(update={"status": "draft"}).model_dump(
                    mode="json"
                ),
                deleted_at=None,
            )

    with pytest.raises(DomainError) as error:
        await OrderService(Repository(), None).export(approved_order.id, "xlsx", "document")
    assert error.value.code == "approval_required"
