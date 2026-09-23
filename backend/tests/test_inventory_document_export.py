import os
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
from uuid import uuid4

import pytest
from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine

from app.core.config import get_settings
from app.core.errors import DomainError
from app.Domains.Catalogs.models import Product, Warehouse
from app.Domains.Integrations1C.models import IntegrationSource
from app.Domains.Inventory.controllers.http import get_service, router
from app.Domains.Inventory.models import Sale
from app.Domains.Inventory.repositories.inventory_repository import InventoryRepository
from app.Domains.Inventory.services import inventory_service
from app.Domains.Inventory.services.inventory_service import InventoryService
from app.Domains.Inventory.services.sales_document_export import XLSX_MEDIA_TYPE
from app.Domains.Security.dependencies import get_current_user


def document_rows(count=2):
    source_id = uuid4()
    return [
        (
            SimpleNamespace(
                id=uuid4(),
                source_id=source_id,
                document_id="=DOCUMENT()",
                date=date(2026, 9, 19),
                line_id=f"=LINE({index})",
                quantity=Decimal("-99999999999999.123456") if index == 0 else Decimal("50.123456"),
                status="posted" if index == 0 else "cancelled",
            ),
            SimpleNamespace(name="=HYPERLINK(1)", sku="+01_23_", code="00123_", unit="@unit"),
            SimpleNamespace(name="-Алматы"),
        )
        for index in range(count)
    ]


def service_for(rows):
    repository = AsyncMock()
    repository.sale.return_value = rows[0][0] if rows else None
    repository.sales_document.return_value = rows
    return InventoryService(repository), repository


async def test_export_preserves_precision_cancelled_rows_ids_and_literal_external_strings():
    rows = document_rows(201)
    service, repository = service_for(rows)
    content = await service.export_sales_document(rows[0][0].id)
    repository.sales_document.assert_awaited_once_with(
        rows[0][0].source_id, rows[0][0].document_id, 100_001
    )
    with BytesIO(content) as stream:
        workbook = load_workbook(stream)
        sheet = workbook["Продажи"]
        assert sheet.max_row == 202  # More than the list endpoint's maximum page size.
        assert sheet.max_column == 7
        assert [cell.value for cell in sheet[1]] == [
            "Артикул",
            "Наименование",
            "Код 1С",
            "Склад",
            "Кол-во",
            "Ед.",
            "Статус",
        ]
        assert sheet["E2"].value == "-99999999999999.123456"
        assert sheet["E3"].value == "50.123456"
        assert sheet["G2"].value == "Проведена"
        assert sheet["G3"].value == "Отменена"
        assert sheet["C2"].value == "00123_"
        assert sheet["B2"].value == "=HYPERLINK(1)"
        assert sheet["A2"].value == "+01_23_"
        assert sheet["D2"].value == "-Алматы"
        assert sheet["F2"].value == "@unit"
        assert rows[0][1].sku == "+01_23_"
        assert rows[0][1].code == "00123_"
        for tab in workbook:
            assert all(cell.data_type != "f" for row in tab for cell in row)
        info = workbook["О выгрузке"]
        assert info["B2"].value == "=DOCUMENT()"
        assert info["B3"].value == str(rows[0][0].source_id)
        assert "Не оригинал накладной" in info["B4"].value
        workbook.close()


async def test_missing_sale_does_not_query_document():
    service, repository = service_for([])
    with pytest.raises(DomainError) as error:
        await service.export_sales_document(uuid4())
    assert (error.value.status_code, error.value.code) == (404, "sale_not_found")
    repository.sales_document.assert_not_awaited()


async def test_oversized_document_rejected_instead_of_truncated(monkeypatch):
    monkeypatch.setattr(inventory_service, "MAX_DOCUMENT_LINES", 1)
    rows = document_rows()
    service, _ = service_for(rows)
    with pytest.raises(DomainError) as error:
        await service.export_sales_document(rows[0][0].id)
    assert (error.value.status_code, error.value.code) == (422, "sales_document_too_large")


async def test_document_repository_filters_both_source_and_document_without_status_filter():
    session = AsyncMock()
    session.execute.return_value = Mock(all=lambda: [])
    source_id = uuid4()
    await InventoryRepository(session).sales_document(source_id, "doc", 100_001)
    query = session.execute.await_args.args[0]
    sql = str(query)
    where = sql.split("WHERE", 1)[1].split("ORDER BY", 1)[0]
    assert "inventory_sales.source_id =" in where
    assert "inventory_sales.document_id =" in where
    assert "warehouse_id" not in where and "product_id" not in where and "status" not in where
    assert source_id in query.compile().params.values()


def application(permissions=None, rows=None):
    app = FastAPI()
    app.include_router(router, prefix="/api/v1")
    service, repository = service_for(document_rows() if rows is None else rows)
    app.dependency_overrides[get_service] = lambda: service
    if permissions is not None:
        app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(
            has_permission=lambda code: code in permissions
        )

    @app.exception_handler(DomainError)
    async def domain_error(request: Request, error: DomainError):
        return JSONResponse(
            status_code=error.status_code, content={"detail": error.detail, "code": error.code}
        )

    return app, repository


@pytest.mark.parametrize(("permissions", "status"), [(None, 401), (set(), 403)])
async def test_export_requires_inventory_read(permissions, status):
    app, repository = application(permissions)
    async with AsyncClient(transport=ASGITransport(app), base_url="http://test") as client:
        response = await client.get(f"/api/v1/inventory/sales/{uuid4()}/document")
    assert response.status_code == status
    repository.sale.assert_not_awaited()


async def test_export_http_binary_headers_and_missing_sale():
    app, repository = application({"inventory.read"})
    sale_id = repository.sale.return_value.id
    async with AsyncClient(transport=ASGITransport(app), base_url="http://test") as client:
        response = await client.get(f"/api/v1/inventory/sales/{sale_id}/document")
        assert response.status_code == 200
        assert response.headers["content-type"] == XLSX_MEDIA_TYPE
        assert response.headers["content-disposition"] == (
            f'attachment; filename="sales-document-{sale_id}.xlsx"'
        )
        assert response.headers["cache-control"] == "private, no-store"
        assert response.content.startswith(b"PK")
        repository.sale.return_value = None
        missing = await client.get(f"/api/v1/inventory/sales/{uuid4()}/document")
        assert missing.status_code == 404
        assert missing.json()["code"] == "sale_not_found"


def test_openapi_documents_binary_auth_headers_and_errors():
    app, _ = application()
    operation = app.openapi()["paths"]["/api/v1/inventory/sales/{sale_id}/document"]["get"]
    assert operation["security"]
    assert set(operation["responses"]) >= {"200", "401", "403", "404", "422"}
    success = operation["responses"]["200"]
    assert set(success["content"]) == {XLSX_MEDIA_TYPE}
    assert success["content"][XLSX_MEDIA_TYPE]["schema"]["format"] == "binary"
    assert set(success["headers"]) == {"Content-Disposition", "Cache-Control"}
    assert operation["parameters"][0]["schema"]["format"] == "uuid"


@pytest.mark.skipif(os.getenv("RUN_DB_TESTS") != "1", reason="opt-in PostgreSQL")
async def test_postgres_export_whole_document_across_products_warehouses_but_not_sources():
    engine = create_async_engine(get_settings().database_url)
    try:
        async with engine.connect() as connection:
            transaction = await connection.begin()
            try:
                async with AsyncSession(connection, expire_on_commit=False) as session:
                    sources = [IntegrationSource(name=f"Export test {index}") for index in range(2)]
                    session.add_all(sources)
                    await session.flush()
                    products, warehouses = [], []
                    for index in range(3):
                        shared = dict(
                            source_id=sources[index // 2].id,
                            source_revision=1,
                            payload_hash="0" * 64,
                            external_id=f"export-{index}",
                        )
                        products.append(
                            Product(**shared, name=f"Товар {index}", sku="001", unit="шт")
                        )
                        warehouses.append(Warehouse(**shared, name=f"Склад {index}"))
                    session.add_all(products + warehouses)
                    await session.flush()
                    sales = []
                    for index in range(202):
                        product_index = index % 2 if index < 201 else 2
                        sales.append(
                            Sale(
                                source_id=products[product_index].source_id,
                                source_revision=1,
                                payload_hash="0" * 64,
                                external_id=f"sale-{index}",
                                product_id=products[product_index].id,
                                warehouse_id=warehouses[product_index].id,
                                date=date(2026, 9, 19),
                                document_id="same-document",
                                line_id=str(index),
                                quantity=Decimal("-5.123456"),
                                status="posted" if index == 0 else "cancelled",
                            )
                        )
                    session.add_all(sales)
                    await session.flush()
                    content = await InventoryService(
                        InventoryRepository(session)
                    ).export_sales_document(sales[0].id)
                    workbook = load_workbook(BytesIO(content))
                    values = list(workbook["Продажи"].values)[1:]
                    assert len(values) == 201
                    assert {row[1] for row in values} == {"Товар 0", "Товар 1"}
                    assert {row[3] for row in values} == {"Склад 0", "Склад 1"}
                    assert {row[4] for row in values} == {"-5.123456"}
                    assert sum(row[6] == "Отменена" for row in values) == 200
                    workbook.close()
            finally:
                await transaction.rollback()
    finally:
        await engine.dispose()
